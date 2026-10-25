"""Workbook upload and validation for the FOU, IAC, WTG, and Programme file."""

from pathlib import Path

import pytest
from apps.datasets.services.parsing import UploadRejected, parse_workbook
from apps.datasets.services.uploads import UploadRejected as StoredRejection
from apps.datasets.services.uploads import activate_dataset, create_dataset_version
from apps.projects.models import Project
from django.contrib.auth import get_user_model
from tests.support import information_workbook, location_row, with_vba

pytestmark = pytest.mark.django_db

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "sequencer_input_v2.xlsx"


def test_supplied_workbook_loads_fou_iac_wtg_and_programme():
    parsed = parse_workbook(FIXTURE.read_bytes())
    assert parsed.error_count == 0
    assert len(parsed.foundations) == 51
    assert len(parsed.cables) == 55
    assert len(parsed.wtg_rows) == 51
    assert len(parsed.programme) == 55
    included = [item for item in parsed.foundations if item.include]
    spiders = [item for item in included if item.attributes.get("Spider location") is True]
    shallows = [item for item in included if item.attributes.get("Shallow location") is True]
    assert len(included) == 50
    assert len(spiders) == 10
    assert len(shallows) == 5
    oss = next(item for item in parsed.foundations if item.foundation_id == "OSS")
    assert oss.include is False
    assert oss.weight_t is None
    assert oss.attributes["Spider location"] is None
    known = {item.foundation_id.casefold() for item in parsed.foundations}
    known.update(item.wtg_id.casefold() for item in parsed.wtg_rows)
    assert all(item.include for item in parsed.cables)
    assert all(item.start_location_id.casefold() in known for item in parsed.cables)
    assert all(item.end_location_id.casefold() in known for item in parsed.cables)
    assert not any(issue.code == "unknown_endpoint" for issue in parsed.issues)
    assert parsed.programme[0].slot == 1
    assert parsed.programme[0].fou_p50_start is not None
    assert parsed.programme[-1].slot == 55
    assert parsed.programme[-1].fou_p50_start is None
    assert parsed.programme[-1].iac_p50_start is not None


def test_valid_workbook_parses_and_can_be_activated():
    user = get_user_model().objects.create_user(username="uploader", password="local-pass-123")
    project = Project.objects.create(owner=user, name="Upload", package_type="Foundation")
    version = create_dataset_version(
        project=project,
        user=user,
        filename="locations.xlsx",
        content=information_workbook([location_row("A"), location_row("B", Location=2)]),
    )
    assert version.error_count == 0
    assert version.row_count == 2
    assert version.foundations.count() == 2
    assert version.cables.count() == 0
    assert version.constraint_rows.count() == 0
    assert version.setting_rows.count() == 0
    activate_dataset(version)
    version.refresh_from_db()
    assert version.is_active is True


def test_missing_sheet_is_a_critical_error():
    content = information_workbook([location_row("A")], omit="IAC data")
    parsed = parse_workbook(content)
    assert parsed.error_count == 1
    assert parsed.issues[0].code == "missing_sheet"


def test_duplicate_foundation_id_is_an_error():
    parsed = parse_workbook(
        information_workbook([location_row("A"), location_row("A", Location=2)])
    )
    assert any(issue.code == "duplicate_foundation_id" for issue in parsed.issues)


def test_invalid_boolean_is_an_error():
    parsed = parse_workbook(
        information_workbook([location_row("A", **{"Spider location": "maybe"})])
    )
    assert any(issue.code == "invalid_boolean" for issue in parsed.issues)


def test_invalid_weight_is_an_error():
    parsed = parse_workbook(information_workbook([location_row("A", **{"Aver. Weight[t]": 0})]))
    assert any(issue.code == "invalid_weight" for issue in parsed.issues)


def test_invalid_dataset_cannot_be_activated():
    user = get_user_model().objects.create_user(username="blocked", password="local-pass-123")
    project = Project.objects.create(owner=user, name="Blocked", package_type="Foundation")
    version = create_dataset_version(
        project=project,
        user=user,
        filename="bad.xlsx",
        content=information_workbook([location_row("A", **{"Aver. Weight[t]": "heavy"})]),
    )
    assert version.error_count > 0
    with pytest.raises(ValueError):
        activate_dataset(version)


def test_macro_workbook_is_rejected_without_execution():
    content = with_vba(information_workbook([location_row("A")]))
    with pytest.raises(UploadRejected):
        parse_workbook(content)
    with pytest.raises(StoredRejection):
        user = get_user_model().objects.create_user(username="macro", password="local-pass-123")
        project = Project.objects.create(owner=user, name="Macro", package_type="Foundation")
        create_dataset_version(project=project, user=user, filename="macro.xlsx", content=content)


def test_formula_without_a_stored_value_is_not_calculated():
    from io import BytesIO

    from openpyxl import load_workbook

    content = information_workbook([location_row("A")])
    workbook = load_workbook(BytesIO(content))
    fou = workbook["FOU data"]
    easting = next(
        cell.column for cell in fou[1] if str(cell.value).casefold().startswith("easting")
    )
    fou.cell(2, easting, "=1+1")
    buffer = BytesIO()
    workbook.save(buffer)
    parsed = parse_workbook(buffer.getvalue())
    assert any(issue.code == "formula" for issue in parsed.issues)
    assert parsed.foundations == [] or parsed.foundations[0].easting_m is None


def test_excluded_row_may_omit_weight_and_unknown_cable_ends_warn():
    parsed = parse_workbook(
        information_workbook(
            [location_row("OSS", Include="NO", **{"Aver. Weight[t]": None})],
            iac_rows=[["C1", "nick", "YES", "A", "AB", "OSS", "MISSING", "NO"]],
            extra_sheet="Notes",
        )
    )
    assert parsed.error_count == 0
    assert parsed.foundations[0].include is False
    assert parsed.foundations[0].weight_t is None
    assert any(issue.code == "unknown_endpoint" for issue in parsed.issues)
