"""Excel export contents."""

from io import BytesIO

import pytest
from apps.constraints.models import Constraint
from apps.datasets.services.uploads import activate_dataset, create_dataset_version
from apps.exports.services.workbook import build_run_workbook
from apps.optimisation.services.runner import execute
from apps.projects.models import Project
from apps.scenarios.models import Scenario, VesselSettings
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from tests.support import information_workbook, location_row

pytestmark = pytest.mark.django_db

REQUIRED_SHEETS = [
    "Optimized_Sequence",
    "Trip_Summary",
    "Validation",
    "Foundation_Input",
    "Vessel_Settings",
    "Constraint_Configuration",
    "Run_Summary",
]


def test_export_contains_the_audit_sheets():
    user = get_user_model().objects.create_user(username="exporter", password="local-pass-123")
    project = Project.objects.create(owner=user, name="Export", package_type="Foundation")
    content = information_workbook(
        [
            location_row("A", **{"Aver. Weight[t]": 100}),
            location_row("B", **{"Aver. Weight[t]": 110, "Location": 2}),
        ]
    )
    version = create_dataset_version(
        project=project, user=user, filename="export.xlsx", content=content
    )
    activate_dataset(version)
    scenario = Scenario.objects.create(project=project, name="Base", dataset_version=version)
    VesselSettings.objects.create(
        scenario=scenario, max_foundations_per_trip=2, solver_time_limit_seconds=10
    )
    Constraint.objects.create(
        scenario=scenario,
        name="A on trip 1",
        constraint_type="fixed_trip",
        enforcement="hard",
        active=True,
        parameters_json={"foundation_id": "A", "trip_number": 1},
        display_order=1,
    )
    run = execute(scenario)
    assert run.validation_status == "Passed"
    workbook = load_workbook(BytesIO(build_run_workbook(run)))
    assert workbook.sheetnames == REQUIRED_SHEETS
    sequence = workbook["Optimized_Sequence"]
    headers = [cell.value for cell in sequence[2]]
    assert headers[0] == "Installation_Sequence"
    data_rows = [row for row in sequence.iter_rows(min_row=3, values_only=True) if row[3]]
    assert len(data_rows) == 2
    assert sequence.freeze_panes == "A3"
    assert sequence.auto_filter.ref.startswith("A2:")

    constraints = workbook["Constraint_Configuration"]
    constraint_text = " ".join(
        str(value) for row in constraints.iter_rows(min_row=1, values_only=True) for value in row
    )
    assert "A on trip 1" in constraint_text
    assert "fixed_trip" in constraint_text
    assert "foundation_id" in constraint_text

    validation = workbook["Validation"]
    validation_text = " ".join(
        str(value) for row in validation.iter_rows(min_row=1, values_only=True) for value in row
    )
    assert "Passed" in validation_text
    summary = " ".join(
        str(value)
        for row in workbook["Run_Summary"].iter_rows(min_row=1, values_only=True)
        for value in row
    )
    assert "Export" in summary
    assert str(run.pk) in summary
