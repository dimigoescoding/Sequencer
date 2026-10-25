"""Create a local demonstration project. Refuses to run when DEBUG is off."""

from io import BytesIO

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from openpyxl import Workbook

from apps.constraints.models import Constraint
from apps.datasets.services.catalogue import compile_constraint
from apps.datasets.services.uploads import activate_dataset, create_dataset_version
from apps.projects.models import Project
from apps.scenarios.models import Scenario, VesselSettings

DEMO_USERNAME = "demo"
DEMO_PASSWORD = "demo-local-only"
DEMO_PROJECT = "Demo Foundation Project"
DEMO_SCENARIO = "Demonstration"


class Command(BaseCommand):
    help = "Load a synthetic local demonstration project. Not for production."

    def handle(self, *args, **options):
        if not settings.DEBUG:
            raise CommandError(
                "seed_demo is for local development only and will not run when DEBUG is off."
            )
        user = get_user_model().objects.filter(username=DEMO_USERNAME).first()
        if user is None:
            user = get_user_model().objects.create_user(
                username=DEMO_USERNAME,
                email="demo@example.com",
                password=DEMO_PASSWORD,
            )
        else:
            user.set_password(DEMO_PASSWORD)
            user.save(update_fields=["password"])
        Project.objects.filter(owner=user, name=DEMO_PROJECT).delete()
        project = Project.objects.create(
            owner=user,
            name=DEMO_PROJECT,
            description="Synthetic foundation campaign for local demonstration. It is not a real wind farm.",
            package_type=Project.PackageType.FOUNDATION,
        )
        version = create_dataset_version(
            project=project,
            user=user,
            filename="synthetic_foundations.xlsx",
            content=_workbook(),
        )
        if version.error_count:
            details = "; ".join(
                version.issues.filter(severity="error").values_list("message", flat=True)[:8]
            )
            raise CommandError(f"Demo workbook failed validation: {details}")
        activate_dataset(version)
        scenario = Scenario.objects.create(
            project=project,
            name=DEMO_SCENARIO,
            description="Synthetic FOU rows. Constraints and trip settings are stored on the scenario.",
            dataset_version=version,
        )
        VesselSettings.objects.create(
            scenario=scenario,
            max_foundations_per_trip=4,
            wtg_per_trip=4,
            solver_time_limit_seconds=20,
        )
        _add_constraints(scenario)
        self.stdout.write(
            self.style.SUCCESS(f"Demo user '{DEMO_USERNAME}' / project '{DEMO_PROJECT}' is ready.")
        )
        self.stdout.write("Local password: demo-local-only")


def _add_constraints(scenario) -> None:
    rows = [
        (
            "FOU",
            "Spider locations FOU",
            "Spider locations should be installed first, but in no particular order.",
            "hard",
            "NA",
            True,
        ),
        (
            "FOU",
            "Shallow locations",
            "Only one shallow location per trip allowed",
            "hard",
            "NA",
            True,
        ),
        (
            "FOU",
            "Infield sailing optimization FOU",
            "Optimizes the infield sailing time within the trip from the given coordinates.",
            "soft",
            "0.5",
            True,
        ),
        (
            "FOU",
            "Jacking difficulty",
            "Avoid difficult locations in the first 3 trips",
            "soft",
            "0.5",
            False,
        ),
    ]
    for index, (package, name, explanation, enforcement, severity, active) in enumerate(
        rows, start=1
    ):
        compiled = compile_constraint(
            package=package,
            name=name,
            explanation=explanation,
            enforcement=enforcement,
            reported_severity=severity,
            active=active,
            source_row=index,
        )
        Constraint.objects.create(
            scenario=scenario,
            name=compiled.name,
            constraint_type=compiled.constraint_type,
            description=compiled.explanation,
            active=compiled.active,
            enforcement=compiled.enforcement,
            severity=compiled.severity,
            package=package,
            parameters_json=compiled.parameters,
            display_order=index,
        )


def _workbook() -> bytes:
    workbook = Workbook()
    fou = workbook.active
    fou.title = "FOU data"
    fou.append(
        [
            "WTG ID",
            "Include",
            "Structure",
            "Easting",
            "Northing",
            "Aver. Weight[t]",
            "Spider location",
            "Shallow location",
            "Jacking Difficulty",
        ]
    )
    rows = [
        ("S1", 0, 0, 500, "YES", "NO", "Easy"),
        ("S2", 100, 0, 520, "YES", "NO", "Easy"),
        ("A1", 0, 100, 480, "NO", "YES", "Easy"),
        ("A2", 100, 100, 490, "NO", "NO", "Difficult"),
        ("B1", 0, 200, 470, "NO", "YES", "Easy"),
        ("B2", 100, 200, 460, "NO", "NO", "Easy"),
        ("C1", 0, 300, 450, "NO", "NO", "Easy"),
        ("C2", 100, 300, 440, "NO", "NO", "Difficult"),
    ]
    for wtg, easting, northing, weight, spider, shallow, difficulty in rows:
        fou.append([wtg, "YES", "WTG", easting, northing, weight, spider, shallow, difficulty])
    iac = workbook.create_sheet("IAC data")
    iac.append(
        [
            "Cable ID",
            "Include",
            "Location ID (Cable start-point)",
            "Location ID (Cable end-point)",
        ]
    )
    wtg_sheet = workbook.create_sheet("WTG data")
    wtg_sheet.append(["WTG ID", "Include", "Easting", "Northing"])
    programme = workbook.create_sheet("Programme")
    programme["B1"] = "FOU"
    programme["D1"] = "IAC"
    programme["F1"] = "WTG"
    programme.merge_cells("B1:C1")
    programme.merge_cells("D1:E1")
    programme.merge_cells("F1:G1")
    for index, name in enumerate(
        ("Location", "P50 Start", "P50 End", "P50 Start", "P50 End", "P50 Start", "P50 End"),
        start=1,
    ):
        programme.cell(2, index, name)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
