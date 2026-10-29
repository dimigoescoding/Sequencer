"""Word report for one successful optimisation run.

A run is successful when the solver status is optimal or feasible and the
independent hard validation passed. Anything else does not produce a file.
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from django.utils.text import slugify
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw

from apps.constraints.library import CONSTRAINT_TYPES
from apps.optimisation.models import OptimisationRun
from apps.optimisation.services.progress import progress_chart

_AMSTERDAM = ZoneInfo("Europe/Amsterdam")
_NAVY = "0F2C4C"
_PACKAGES = ("FOU", "IAC", "WTG")
_SETTINGS = (
    ("max_foundations_per_trip", "FOU per trip"),
    ("wtg_per_trip", "WTG per trip"),
    ("solver_time_limit_seconds", "Solver time limit (seconds)"),
    ("allow_partial_final_trip", "Partial final trip"),
    ("base_maximum_trip_payload_t", "Base maximum trip payload (t)"),
    ("transit_speed_knots", "Transit speed (knots)"),
    ("harbour_easting_m", "Harbour easting (m)"),
    ("harbour_northing_m", "Harbour northing (m)"),
)
_ABBREVIATIONS = (
    ("FOU", "Foundation installation package"),
    ("IAC", "Inter-array cable package"),
    ("WTG", "Wind turbine generator package"),
    ("MP", "Monopile, the foundation installation activity on the progress graph"),
    ("SSt", "Secondary steel"),
    ("P50", "Programme date used for an activity"),
    ("CP-SAT", "OR-Tools constraint-programming solver"),
)
_PROGRAMME = (
    ("FOU", "fou_p50_start", "fou_p50_end"),
    ("SSt", "sst_p50_start", "sst_p50_end"),
    ("Filter", "filter_p50_start", "filter_p50_end"),
    ("Armour", "armour_p50_start", "armour_p50_end"),
    ("IAC lay", "iac_p50_start", "iac_p50_end"),
    ("IAC burial", "burial_p50_start", "burial_p50_end"),
    ("WTG", "wtg_p50_start", "wtg_p50_end"),
    ("Commissioning", "commissioning_p50_start", "commissioning_p50_end"),
)


class ReportNotAvailable(Exception):
    """The run is not successful, so no report file is built."""


def report_filename(run: OptimisationRun) -> str:
    project = slugify(run.scenario.project.name)[:40] or "project"
    scenario = slugify(run.scenario.name)[:40] or "scenario"
    return f"sequencer_{project}_{scenario}_run{run.pk}.docx"


def build_run_report(run: OptimisationRun) -> bytes:
    if not run.approved:
        raise ReportNotAvailable("A report is available only for a successful run.")
    document = Document()
    _page_setup(document)
    _footer(document, run)
    _front_page(document, run)
    document.add_page_break()
    _contents(document)
    _abbreviations(document)
    _introduction(document, run)
    _settings(document, run)
    _constraints(document, run)
    _results(document, run)
    _blank_section(
        document,
        "Risk and opportunities",
        "To be completed after this report is downloaded.",
    )
    _blank_section(
        document,
        "Conclusions and recommendations",
        "To be completed after this report is downloaded.",
    )
    _appendix_sequence(document, run)
    _appendix_dataset(document, run)
    _appendix_solver(document)
    _update_fields_on_open(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _page_setup(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x1C, 0x28, 0x33)


def _front_page(document: Document, run: OptimisationRun) -> None:
    document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("Sequencer")
    run_title.bold = True
    run_title.font.size = Pt(28)
    run_title.font.color.rgb = RGBColor(0x0F, 0x2C, 0x4C)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Installation sequence report")
    subtitle_run.font.size = Pt(18)
    subtitle_run.font.color.rgb = RGBColor(0x0F, 0x2C, 0x4C)
    document.add_paragraph()
    for label, value in (
        ("Project", run.scenario.project.name),
        ("Scenario", run.scenario.name),
        ("Dataset", f"Version {run.dataset_version.version_number}"),
        ("Run", str(run.pk)),
        ("Date", _report_date(run.created_at)),
        ("Result", f"{run.solver_status}, validation {run.validation_status}"),
    ):
        line = document.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_run = line.add_run(f"{label}  ")
        label_run.bold = True
        line.add_run(value)


def _contents(document: Document) -> None:
    heading = document.add_paragraph()
    heading_run = heading.add_run("Table of contents")
    heading_run.bold = True
    heading_run.font.size = Pt(16)
    heading_run.font.color.rgb = RGBColor(0x0F, 0x2C, 0x4C)
    paragraph = document.add_paragraph()
    _field(paragraph, r'TOC \o "1-2" \h \z \u')
    note = document.add_paragraph()
    note_run = note.add_run("Word fills this list from the headings when the file is opened.")
    note_run.italic = True
    note_run.font.size = Pt(10)


def _abbreviations(document: Document) -> None:
    document.add_heading("Abbreviations", level=1)
    _table(document, ["Abbreviation", "Meaning"], [list(item) for item in _ABBREVIATIONS])


def _introduction(document: Document, run: OptimisationRun) -> None:
    document.add_heading("Introduction", level=1)
    packages = _packages(run)
    if packages:
        package_sentence = f"The packages in this sequence are {_join(packages)}."
    else:
        package_sentence = "No package rows were stored in this sequence."
    document.add_paragraph(
        f"{run.scenario.project.name} is sequenced in Sequencer. "
        f"This report covers scenario {run.scenario.name} and dataset version "
        f"{run.dataset_version.version_number}. "
        f"{package_sentence}"
    )


def _settings(document: Document, run: OptimisationRun) -> None:
    document.add_heading("Settings", level=1)
    document.add_paragraph(
        "These are the settings stored with the run. A blank payload, speed, or harbour "
        "value is not given an engineering default."
    )
    snapshot = run.vessel_snapshot or {}
    rows = [[label, _setting_value(snapshot.get(key))] for key, label in _SETTINGS]
    _table(document, ["Setting", "Value"], rows)


def _constraints(document: Document, run: OptimisationRun) -> None:
    document.add_heading("Constraints", level=1)
    document.add_paragraph("Active constraints only. An inactive constraint was not sent to the solver.")
    rows = []
    items = [
        item
        for item in (run.constraint_snapshot or [])
        if item.get("active")
    ]
    items.sort(key=lambda item: (item.get("display_order") or 0, str(item.get("name") or "")))
    for item in items:
        kind = str(item.get("constraint_type") or "")
        label = CONSTRAINT_TYPES.get(kind, {}).get("label", kind or "—")
        enforcement = str(item.get("enforcement") or "").capitalize() or "—"
        severity = item.get("severity")
        rows.append(
            [
                item.get("name") or "—",
                item.get("description") or "—",
                f"{label} ({enforcement})",
                "—" if severity in (None, "") else str(severity),
            ]
        )
    if rows:
        _table(document, ["Name", "Description", "Type", "Severity"], rows, size=9)
    else:
        document.add_paragraph("No active constraints were stored with this run.")


def _results(document: Document, run: OptimisationRun) -> None:
    document.add_heading("Results", level=1)
    document.add_paragraph(
        "The installation sequence for each package is in Appendix A. "
        f"The solver status is {run.solver_status} and validation is {run.validation_status}."
    )
    chart = progress_chart(run)
    if chart is None:
        document.add_paragraph(
            "No programme end dates are stored for this sequence, so the progress graph is not included."
        )
        return
    image = BytesIO(_chart_png(chart))
    document.add_picture(image, width=Cm(16.5))
    caption = document.add_paragraph()
    caption_run = caption.add_run(
        "Cumulative number installed against time. Sequence position 1 uses programme slot 1. "
        "A location is counted when that activity ends."
    )
    caption_run.italic = True
    caption_run.font.size = Pt(10)


def _blank_section(document: Document, title: str, text: str) -> None:
    document.add_heading(title, level=1)
    document.add_paragraph(text)


def _appendix_sequence(document: Document, run: OptimisationRun) -> None:
    document.add_page_break()
    document.add_heading("Appendix A — Installation sequence", level=1)
    document.add_paragraph("Stored order for each package. Trip and vessel position follow from that order.")
    grouped = {package: [] for package in _PACKAGES}
    for row in run.sequence_rows.all():
        package = str(row.package or "FOU").strip().upper()
        if package not in grouped:
            package = "FOU"
        grouped[package].append(row)
    written = False
    for package in _PACKAGES:
        rows = sorted(grouped[package], key=lambda item: item.installation_sequence)
        if not rows:
            continue
        written = True
        document.add_heading(package, level=2)
        _table(
            document,
            ["Sequence", "Trip", "Vessel position", "ID", "Weight (t)"],
            [
                [
                    row.installation_sequence,
                    row.trip,
                    row.vessel_position,
                    row.foundation_id,
                    _plain(row.weight_t),
                ]
                for row in rows
            ],
            size=9,
        )
    if not written:
        document.add_paragraph("No sequence rows were stored.")


def _appendix_dataset(document: Document, run: OptimisationRun) -> None:
    document.add_page_break()
    document.add_heading("Appendix B — Dataset", level=1)
    dataset = run.dataset_version
    filename = dataset.original_filename or "uploaded workbook"
    document.add_paragraph(
        f"Dataset version {dataset.version_number} ({filename}). "
        "Include marks the rows that were available to the solver."
    )
    document.add_heading("FOU data", level=2)
    _table(
        document,
        ["ID", "Include", "Weight (t)", "Easting", "Northing"],
        [
            [
                row.foundation_id,
                _yes_no(row.include),
                _plain(row.weight_t),
                _plain(row.easting_m),
                _plain(row.northing_m),
            ]
            for row in dataset.foundations.all()
        ],
        size=8,
    )
    document.add_heading("IAC data", level=2)
    _table(
        document,
        ["Cable ID", "Include", "String ID", "Loop ID", "Start", "End", "Spider"],
        [
            [
                row.cable_id,
                _yes_no(row.include),
                row.string_id or "—",
                row.loop_id or "—",
                row.start_location_id or "—",
                row.end_location_id or "—",
                _yes_no(row.spider_cable),
            ]
            for row in dataset.cables.all()
        ],
        size=8,
    )
    document.add_heading("WTG data", level=2)
    _table(
        document,
        ["ID", "Include", "Structure", "Easting", "Northing", "String ID"],
        [
            [
                row.wtg_id,
                _yes_no(row.include),
                row.structure or "—",
                _plain(row.easting_m),
                _plain(row.northing_m),
                row.string_label or "—",
            ]
            for row in dataset.wtg_locations.all()
        ],
        size=8,
    )
    document.add_heading("Programme", level=2)
    document.add_paragraph("Each cell is the P50 start and end. Times are UTC.")
    _table(
        document,
        ["Slot", *[name for name, _start, _end in _PROGRAMME]],
        [
            [slot.slot, *[_window(slot, start, end) for _name, start, end in _PROGRAMME]]
            for slot in dataset.programme_slots.all()
        ],
        size=7,
    )


def _appendix_solver(document: Document) -> None:
    document.add_page_break()
    document.add_heading("Appendix C — Solver method", level=1)
    document.add_paragraph(
        "Sequencer uses Google OR-Tools CP-SAT. Each package has one installation order. "
        "The trip number and the position on the vessel follow from that order and the trip size. "
        "When IAC or WTG locations are included, FOU, IAC, and WTG are solved together in one model. "
        "The packages do not wait for each other unless a Package buffer constraint says so."
    )
    document.add_paragraph(
        "A hard constraint is a requirement. The solver has to meet it. A hard constraint is not "
        "turned into a soft one, and it is not dropped to produce a sequence. After the solver stops, "
        "the same hard rules are checked again on the stored order. If that check fails, the run is "
        "invalid and is not an approved sequence. This report is produced only when that check passed."
    )
    document.add_paragraph(
        "A soft constraint is a preference. Severity from 1 to 100 sets how strongly it pulls against "
        "the other soft preferences. A soft break keeps the run valid and records a penalty."
    )
    document.add_paragraph(
        "Programme slot 1 supplies the dates for sequence position 1, slot 2 for position 2, and so on. "
        "SSt, filter layer, and armour layer follow the FOU order. IAC burial follows the IAC order. "
        "WTG commissioning follows the WTG order. Minimise total sailing distance shortens the package "
        "it is assigned to. The distance is the straight line from one item to the next, including the "
        "join from one trip to the next. Harbour transit is not included."
    )
    document.add_paragraph(
        "The search uses one worker and a fixed seed, so the same inputs follow the same search. "
        "The time limit is 1 to 600 seconds. If no feasible order exists, the search ends without a "
        "proven result, or the solver reports an error, no sequence is stored."
    )


def _table(document: Document, headers: list[str], rows: list[list], size: int = 10) -> None:
    if not rows:
        document.add_paragraph("No rows stored.")
        return
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        _cell(table.rows[0].cells[index], header, size=size, bold=True, white=True)
    _shade(table.rows[0], _NAVY)
    _repeat_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            _cell(cells[index], value, size=size)
    document.add_paragraph()


def _cell(cell, value, *, size: int, bold: bool = False, white: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(_plain(value))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if white:
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _shade(row, color: str) -> None:
    for cell in row.cells:
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color)
        shading.set(qn("w:val"), "clear")
        properties.append(shading)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    properties.append(marker)


def _field(paragraph, instruction: str) -> None:
    begin = paragraph.add_run()
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    begin._r.append(start)
    code = paragraph.add_run()
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    code._r.append(text)
    separate = paragraph.add_run()
    middle = OxmlElement("w:fldChar")
    middle.set(qn("w:fldCharType"), "separate")
    separate._r.append(middle)
    paragraph.add_run(" ")
    end = paragraph.add_run()
    stop = OxmlElement("w:fldChar")
    stop.set(qn("w:fldCharType"), "end")
    end._r.append(stop)


def _update_fields_on_open(document: Document) -> None:
    settings = document.settings.element
    flag = OxmlElement("w:updateFields")
    flag.set(qn("w:val"), "true")
    settings.append(flag)


def _footer(document: Document, run: OptimisationRun) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run(f"{run.scenario.project.name} · Run {run.pk} · ")
    _field(paragraph, "PAGE")


def _chart_png(chart: dict) -> bytes:
    series = list(chart["series"])
    legend_rows = max(1, (len(series) + 2) // 3)
    width = int(chart["width"])
    height = int(chart["height"]) + 16 + 18 * legend_rows
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    shift = 24
    draw.text((chart["time_x"] - 90, 4), "Number installed", fill="#0F2C4C")
    left = chart["axis_left"]
    top = chart["axis_top"] + shift
    bottom = chart["axis_bottom"] + shift
    right = chart["axis_right"]
    draw.line((left, top, left, bottom), fill="#8AA0B5", width=1)
    draw.line((left, bottom, right, bottom), fill="#8AA0B5", width=1)
    for tick in chart["yticks"]:
        y = float(tick["y"]) + shift
        draw.line((left - 4, y, left, y), fill="#8AA0B5")
        draw.text((8, y - 6), tick["label"], fill="#52606D")
    for tick in chart["xticks"]:
        x = float(tick["x"])
        draw.text((x - 28, bottom + 6), tick["label"], fill="#52606D")
    draw.text((chart["time_x"] - 12, bottom + 22), "Time", fill="#0F2C4C")
    for item in series:
        points = []
        for pair in item["points"].split():
            x_text, y_text = pair.split(",")
            points.append((float(x_text), float(y_text) + shift))
        if len(points) >= 2:
            draw.line(points, fill=item["color"], width=3)
    for index, item in enumerate(series):
        column = index % 3
        row = index // 3
        origin_x = 16 + column * 310
        origin_y = int(chart["height"]) + 4 + row * 18
        draw.rectangle((origin_x, origin_y + 2, origin_x + 14, origin_y + 12), fill=item["color"])
        draw.text((origin_x + 18, origin_y), f"{item['name']} {item['total']}", fill="#1C2833")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _packages(run: OptimisationRun) -> list[str]:
    found = set()
    for row in run.sequence_rows.all():
        package = str(row.package or "FOU").strip().upper()
        if package in _PACKAGES:
            found.add(package)
    return [package for package in _PACKAGES if package in found]


def _join(items: list[str]) -> str:
    if len(items) == 1:
        return items[0]
    if len(items) == 2:
        return f"{items[0]} and {items[1]}"
    return f"{', '.join(items[:-1])}, and {items[-1]}"


def _setting_value(value) -> str:
    if value in (None, ""):
        return "Not set"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def _plain(value) -> str:
    if value is None or value == "":
        return "—"
    return str(value)


def _yes_no(value) -> str:
    if value is True:
        return "Yes"
    if value is False:
        return "No"
    return "—"


def _window(slot, start_name: str, end_name: str) -> str:
    start = _when(getattr(slot, start_name, None))
    end = _when(getattr(slot, end_name, None))
    if start == "—" and end == "—":
        return "—"
    return f"{start} – {end}"


def _when(value) -> str:
    if not isinstance(value, datetime):
        return "—"
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    value = value.astimezone(UTC)
    if value.hour == value.minute == value.second == 0:
        return value.strftime("%Y-%m-%d")
    return value.strftime("%Y-%m-%d %H:%M")


def _report_date(value: datetime | None) -> str:
    if value is None:
        return "—"
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(_AMSTERDAM).strftime("%Y-%m-%d %H:%M Europe/Amsterdam")
