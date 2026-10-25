"""Build an auditable workbook for one optimisation run."""

from __future__ import annotations

import json
from io import BytesIO

from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.optimisation.models import OptimisationRun
from apps.optimisation.services.routing import nautical_miles

HEADER_FILL = PatternFill("solid", fgColor="0F2C4C")
HEADER_FONT = Font(color="FFFFFF", bold=True)
META_FILL = PatternFill("solid", fgColor="E8EEF4")
HARD_FILL = PatternFill("solid", fgColor="F8D7DA")
SOFT_FILL = PatternFill("solid", fgColor="D6EAF8")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)

SEQUENCE_COLUMNS = [
    "Installation_Sequence",
    "Package",
    "Trip",
    "Vessel_Position",
    "Foundation_ID",
    "Weight_t",
    "Trip_Total_Weight_t",
    "Delivery_Sequence",
    "Sequence_Deviation",
    "Delivery_Batch",
    "Difficulty_Value",
    "Punch_Through",
    "Fixed_Trip",
    "Fixed_Position",
    "Earliest_Sequence",
    "Latest_Sequence",
    "Hard_Constraint_Status",
    "Soft_Penalty",
    "Governing_Constraints",
    "Notes",
]


def export_filename(run: OptimisationRun) -> str:
    project = slugify(run.scenario.project.name)[:40] or "project"
    scenario = slugify(run.scenario.name)[:40] or "scenario"
    return f"sequencer_{project}_{scenario}_run{run.pk}.xlsx"


def build_run_workbook(run: OptimisationRun) -> bytes:
    workbook = Workbook()
    sequence = workbook.active
    sequence.title = "Optimized_Sequence"
    _write_table(
        sequence, run, SEQUENCE_COLUMNS, [_sequence_values(row) for row in run.sequence_rows.all()]
    )
    _write_table(
        workbook.create_sheet("Trip_Summary"),
        run,
        [
            "Package",
            "Trip",
            "Foundation_Count",
            "Foundations_In_Order",
            "Payload_Limit_t",
            "Total_Weight_t",
            "Remaining_Payload_t",
            "Punch_Through_Count",
            "Average_Difficulty",
            "Maximum_Difficulty",
            "Route_Distance_nm",
            "Payload_Check",
            "Position_Check",
            "Overall_Trip_Check",
        ],
        [_trip_values(row) for row in run.trip_rows.all()],
    )
    _write_table(
        workbook.create_sheet("Validation"),
        run,
        [
            "Validation_Check",
            "Constraint_Name",
            "Target",
            "Actual_Result",
            "Status",
            "Affected_Foundations",
            "Explanation",
        ],
        [
            [
                row.check_name,
                row.constraint_name,
                row.target,
                row.actual_result,
                row.status,
                row.affected_foundations,
                row.explanation,
            ]
            for row in run.validation_rows.all()
        ],
    )
    _write_table(
        workbook.create_sheet("Foundation_Input"),
        run,
        [
            "Foundation_ID",
            "Include",
            "Weight_t",
            "Delivery_Sequence",
            "Length_m",
            "Delivery_Batch",
            "Easting_m",
            "Northing_m",
            "Punch_Through",
            "Difficulty_Value",
            "Eccentricity_Limit_m",
            "Fixed_Trip",
            "Fixed_Position",
            "Earliest_Sequence",
            "Latest_Sequence",
            "Notes",
        ],
        [_foundation_values(row) for row in run.dataset_version.foundations.all()],
    )
    _write_settings(workbook.create_sheet("Vessel_Settings"), run)
    _write_constraints(workbook.create_sheet("Constraint_Configuration"), run)
    _write_summary(workbook.create_sheet("Run_Summary"), run)
    workbook.properties.title = f"Sequencer run {run.pk}"
    workbook.properties.subject = f"{run.scenario.project.name} / {run.scenario.name}"
    workbook.properties.creator = "Sequencer"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_table(
    worksheet: Worksheet, run: OptimisationRun, headers: list[str], rows: list[list]
) -> None:
    _meta_row(worksheet, run)
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(2, column, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN
    for row_index, values in enumerate(rows, start=3):
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row_index, column, _cell(value))
            cell.border = THIN
            cell.alignment = Alignment(vertical="center")
            _apply_number_format(cell, headers[column - 1], value)
    last_row = max(2, 2 + len(rows))
    last_column = get_column_letter(len(headers))
    worksheet.auto_filter.ref = f"A2:{last_column}{last_row}"
    worksheet.freeze_panes = "A3"
    worksheet.auto_filter.ref = f"A2:{last_column}{max(last_row, 2)}"
    _widths(worksheet, headers, rows)
    worksheet.oddHeader.left.text = f"{run.scenario.project.name} / {run.scenario.name}"
    worksheet.oddFooter.left.text = f"Run {run.pk} | {run.solver_status} | {run.validation_status}"


def _meta_row(worksheet: Worksheet, run: OptimisationRun) -> None:
    values = [
        "Project",
        run.scenario.project.name,
        "Scenario",
        run.scenario.name,
        "Run_ID",
        run.pk,
        "Solver_Status",
        run.solver_status,
        "Validation_Status",
        run.validation_status,
    ]
    for column, value in enumerate(values, start=1):
        cell = worksheet.cell(1, column, value)
        cell.fill = META_FILL
        cell.font = Font(bold=column % 2 == 1)


def _write_settings(worksheet: Worksheet, run: OptimisationRun) -> None:
    _meta_row(worksheet, run)
    headers = ["Setting", "Value"]
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(2, column, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    snapshot = run.vessel_snapshot or {}
    rows = [[key, "Not set" if value in (None, "") else value] for key, value in snapshot.items()]
    rows.append(
        [
            "note",
            "Blank payload, speed, and harbour values are not given engineering defaults.",
        ]
    )
    for row_index, values in enumerate(rows, start=3):
        for column, value in enumerate(values, start=1):
            worksheet.cell(row_index, column, _cell(value)).border = THIN
    worksheet.freeze_panes = "A3"
    worksheet.auto_filter.ref = f"A2:B{2 + len(rows)}"
    worksheet.column_dimensions["A"].width = 36
    worksheet.column_dimensions["B"].width = 88


def _write_constraints(worksheet: Worksheet, run: OptimisationRun) -> None:
    headers = [
        "Name",
        "Constraint_Type",
        "Description",
        "Active",
        "Enforcement",
        "Severity",
        "Scope",
        "Display_Order",
        "Parameters",
    ]
    rows = []
    for item in run.constraint_snapshot or []:
        rows.append(
            [
                item.get("name"),
                item.get("constraint_type"),
                item.get("description") or "",
                item.get("active"),
                item.get("enforcement"),
                item.get("severity"),
                item.get("scope") or "",
                item.get("display_order"),
                json.dumps(item.get("parameters") or {}, sort_keys=True),
            ]
        )
    _write_table(worksheet, run, headers, rows)
    for row_index in range(3, 3 + len(rows)):
        enforcement = worksheet.cell(row_index, 5).value
        fill = HARD_FILL if enforcement == "hard" else SOFT_FILL if enforcement == "soft" else None
        if fill is not None:
            worksheet.cell(row_index, 5).fill = fill


def _write_summary(worksheet: Worksheet, run: OptimisationRun) -> None:
    headers = ["Parameter", "Value"]
    dataset = run.dataset_version
    rows = [
        ["Project", run.scenario.project.name],
        ["Scenario", run.scenario.name],
        ["Package", run.scenario.project.package_type],
        ["Dataset_Version", dataset.version_number],
        ["Dataset_SHA256", dataset.content_sha256],
        ["Run_ID", run.pk],
        ["Run_Date", run.created_at.isoformat() if run.created_at else ""],
        ["Solver_Status", run.solver_status],
        ["Run_Status", run.status],
        ["Validation_Status", run.validation_status],
        ["Runtime_Seconds", run.runtime_seconds],
        ["Objective_Value", run.objective_value],
        ["Objective_Note", "Scaled integer. It is not tonnes, metres, or currency."],
        ["Random_Seed", run.seed],
        ["Included_Foundations", run.included_count],
        ["Trips", run.trip_count],
        ["Active_Hard_Constraints", run.hard_constraint_count],
        ["Active_Soft_Constraints", run.soft_constraint_count],
        ["Soft_Tradeoffs", run.soft_tradeoff_count],
        ["Message", run.message],
        [
            "Penalty_Scaling",
            "Soft penalties use ceil(raw * severity * 1000 / max_raw). See docs/OPTIMISATION_METHOD.md.",
        ],
        ["Approved_Solution", "Yes" if run.approved else "No"],
    ]
    _write_table(worksheet, run, headers, rows)
    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 100


def _sequence_values(row) -> list:
    return [
        row.installation_sequence,
        row.package,
        row.trip,
        row.vessel_position,
        row.foundation_id,
        row.weight_t,
        row.trip_total_weight_t,
        row.delivery_sequence,
        row.sequence_deviation,
        row.delivery_batch,
        row.difficulty_value,
        _bool_text(row.punch_through),
        row.fixed_trip,
        row.fixed_position,
        row.earliest_sequence,
        row.latest_sequence,
        row.hard_constraint_status,
        row.soft_penalty,
        row.governing_constraints,
        row.notes,
    ]


def _trip_values(row) -> list:
    return [
        row.package,
        row.trip,
        row.foundation_count,
        row.foundations_in_order,
        row.payload_limit_t,
        row.total_weight_t,
        row.remaining_payload_t,
        row.punch_through_count,
        row.average_difficulty,
        row.maximum_difficulty,
        None if row.route_distance is None else nautical_miles(row.route_distance),
        row.payload_check,
        row.position_check,
        row.overall_trip_check,
    ]


def _foundation_values(row) -> list:
    return [
        row.foundation_id,
        _bool_text(row.include),
        row.weight_t,
        row.delivery_sequence,
        row.length_m,
        row.delivery_batch,
        row.easting_m,
        row.northing_m,
        _bool_text(row.punch_through),
        row.difficulty_value,
        row.eccentricity_limit_m,
        row.fixed_trip,
        row.fixed_position,
        row.earliest_sequence,
        row.latest_sequence,
        row.notes,
    ]


def _bool_text(value) -> str:
    if value is None:
        return ""
    return "Yes" if value else "No"


def _cell(value):
    if hasattr(value, "as_tuple"):
        return float(value)
    return value


def _apply_number_format(cell, header: str, value) -> None:
    if value is None or isinstance(value, str):
        return
    if header.endswith("_t") or header in {
        "Weight_t",
        "Average_Difficulty",
        "Maximum_Difficulty",
        "Length_m",
    }:
        cell.number_format = "0.000"
    elif header == "Route_Distance_nm":
        cell.number_format = "0.000"
    elif isinstance(value, float):
        cell.number_format = "0.000"


def _widths(worksheet: Worksheet, headers: list[str], rows: list[list]) -> None:
    for index, header in enumerate(headers, start=1):
        width = len(header) + 2
        for values in rows[:40]:
            if index - 1 < len(values) and values[index - 1] is not None:
                width = max(width, min(42, len(str(values[index - 1])) + 2))
        worksheet.column_dimensions[get_column_letter(index)].width = width
