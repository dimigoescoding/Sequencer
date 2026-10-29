"""Parse the FOU, IAC, WTG, and Programme workbook.

Spreadsheet cells are untrusted data. Formula text is never calculated. When
Excel has already stored a calculated value, that stored value is read.
Other sheets are ignored. Constraints and settings are not read from the file.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from apps.optimisation.services.types import FoundationInput

TRUE_TEXT = {"yes", "true", "1", "y"}
FALSE_TEXT = {"no", "false", "0", "n"}

FOU_FIELDS = {
    "wtg id": "wtg_id",
    "include": "include",
    "structure": "structure",
    "easting": "easting",
    "northing": "northing",
    "aver weight": "weight",
    "length": "length",
    "delivery batch": "delivery_batch",
    "spider location": "spider_location",
    "shallow location": "shallow_location",
    "soil difficulty": "soil_difficulty",
    "jacking difficulty": "jacking_difficulty",
    "punch through": "punch_through",
    "eccentricity limit": "eccentricity_limit",
}
IAC_FIELDS = {
    "cable id": "cable_id",
    "nickname": "nickname",
    "include": "include",
    "string id": "string_id",
    "loop id": "loop_id",
    "location id (cable start-point)": "start_location",
    "location id (cable end-point)": "end_location",
    "spider cable": "spider_cable",
}
WTG_FIELDS = {
    "wtg id": "wtg_id",
    "include": "include",
    "structure": "structure",
    "easting": "easting",
    "northing": "northing",
    "spider location": "spider_location",
    "string": "string",
    "string id": "string",
    "jacking difficulty": "jacking_difficulty",
    "jacking difficulty wtg": "jacking_difficulty",
}
FOU_REQUIRED = ("wtg_id", "include", "weight", "easting", "northing")
IAC_REQUIRED = ("cable_id", "include", "start_location", "end_location")
WTG_REQUIRED = ("wtg_id", "include", "easting", "northing")
FIELD_LABELS = {
    "wtg_id": "WTG ID",
    "include": "Include",
    "weight": "Aver. Weight[t]",
    "easting": "Easting",
    "northing": "Northing",
    "cable_id": "Cable ID",
    "start_location": "Location ID (Cable start-point)",
    "end_location": "Location ID (Cable end-point)",
}
SOLVER_COLUMNS = {
    "wtg_id": "Foundation_ID",
    "include": "Include",
    "weight": "Weight_t",
    "easting": "Easting_m",
    "northing": "Northing_m",
    "length": "Length_m",
    "delivery_batch": "Delivery_Batch",
    "spider_location": "Spider location",
    "shallow_location": "Shallow location",
    "jacking_difficulty": "Jacking difficulty",
    "structure": "Structure",
    "soil_difficulty": "Soil difficulty",
    "punch_through": "Punch_Through",
    "eccentricity_limit": "Eccentricity_Limit_m",
}
IAC_COLUMNS = {
    "cable_id": "Cable ID",
    "nickname": "Nickname",
    "string_id": "String ID",
    "loop_id": "Loop ID",
    "start_location": "Start location",
    "end_location": "End location",
    "spider_cable": "Spider cable",
}
WTG_COLUMNS = {
    "wtg_id": "Foundation_ID",
    "include": "Include",
    "structure": "Structure",
    "easting": "Easting_m",
    "northing": "Northing_m",
    "spider_location": "Spider location",
    "string": "String ID",
    "jacking_difficulty": "Jacking difficulty WTG",
}
REQUIRED_SHEETS = ("FOU data", "IAC data", "WTG data", "Programme")


class UploadRejected(Exception):
    """The file cannot be stored as a dataset version."""


@dataclass
class Issue:
    severity: str
    code: str
    message: str
    row_number: int | None = None
    column_name: str = ""


@dataclass
class CableRow:
    cable_id: str
    include: bool | None
    nickname: str
    string_id: str
    loop_id: str
    start_location_id: str
    end_location_id: str
    spider_cable: bool | None
    source_row: int


@dataclass
class WtgRow:
    wtg_id: str
    include: bool | None
    structure: str
    easting_m: Decimal | None
    northing_m: Decimal | None
    spider_location: bool | None
    string_label: str
    jacking_difficulty: str
    source_row: int


_ACTIVITY_HEADERS = {
    "fou": "fou",
    "fou installation": "fou",
    "sst": "sst",
    "sst installation": "sst",
    "s st installation": "sst",
    "filter layer": "filter",
    "scour filter": "filter",
    "armour layer": "armour",
    "armor layer": "armour",
    "scour armour": "armour",
    "scour armor": "armour",
    "iac": "iac",
    "iac lay": "iac",
    "iac burial": "burial",
    "wtg": "wtg",
    "wtg installation": "wtg",
    "wtg commissioning": "commissioning",
    "wtg commisioning": "commissioning",
    "wtg commisoning": "commissioning",
    "wtg commissioned": "commissioning",
}
_REQUIRED_ACTIVITIES = ("fou", "iac", "wtg")
_ACTIVITY_NAMES = {
    "fou": "FOU installation",
    "sst": "SSt installation",
    "filter": "Filter layer",
    "armour": "Armour layer",
    "iac": "IAC lay",
    "burial": "IAC burial",
    "wtg": "WTG installation",
    "commissioning": "WTG commissioning",
}


@dataclass
class ProgrammeRow:
    slot: int
    source_row: int
    fou_p50_start: datetime | None = None
    fou_p50_end: datetime | None = None
    sst_p50_start: datetime | None = None
    sst_p50_end: datetime | None = None
    filter_p50_start: datetime | None = None
    filter_p50_end: datetime | None = None
    armour_p50_start: datetime | None = None
    armour_p50_end: datetime | None = None
    iac_p50_start: datetime | None = None
    iac_p50_end: datetime | None = None
    burial_p50_start: datetime | None = None
    burial_p50_end: datetime | None = None
    wtg_p50_start: datetime | None = None
    wtg_p50_end: datetime | None = None
    commissioning_p50_start: datetime | None = None
    commissioning_p50_end: datetime | None = None


@dataclass
class ParseResult:
    foundations: list[FoundationInput] = field(default_factory=list)
    cables: list[CableRow] = field(default_factory=list)
    wtg_rows: list[WtgRow] = field(default_factory=list)
    programme: list[ProgrammeRow] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    columns: list[str] = field(default_factory=list)
    source_rows: list[int] = field(default_factory=list)

    @property
    def error_count(self) -> int:
        return sum(1 for issue in self.issues if issue.severity == "error")

    @property
    def warning_count(self) -> int:
        return sum(1 for issue in self.issues if issue.severity == "warning")


def parse_workbook(content: bytes) -> ParseResult:
    if _contains_macros(content):
        raise UploadRejected("Macro-enabled workbooks are not accepted and their code is not run.")
    try:
        formulas = load_workbook(BytesIO(content), data_only=False, keep_vba=False)
        cached = load_workbook(BytesIO(content), data_only=True, keep_vba=False)
    except Exception as exc:
        raise UploadRejected("The file is not a readable .xlsx workbook.") from exc

    sheets = {label: _find_sheet(formulas, label) for label in REQUIRED_SHEETS}
    missing = [label for label, sheet in sheets.items() if sheet is None]
    if missing:
        return ParseResult(
            issues=[
                Issue(
                    "error",
                    "missing_sheet",
                    "The workbook must contain FOU data, IAC data, WTG data, and Programme sheets. "
                    "Missing: " + ", ".join(missing) + ".",
                )
            ]
        )

    issues: list[Issue] = []
    foundations, source_rows, fou_columns, fou_issues = _parse_table(
        sheets["FOU data"],
        _find_sheet(cached, "FOU data"),
        FOU_FIELDS,
        FOU_REQUIRED,
        _foundation_from_values,
        SOLVER_COLUMNS,
    )
    issues.extend(fou_issues)
    cables, _cable_rows, iac_columns, iac_issues = _parse_table(
        sheets["IAC data"],
        _find_sheet(cached, "IAC data"),
        IAC_FIELDS,
        IAC_REQUIRED,
        _cable_from_values,
        IAC_COLUMNS,
    )
    issues.extend(iac_issues)
    wtg_rows, _wtg_source, wtg_columns, wtg_issues = _parse_table(
        sheets["WTG data"],
        _find_sheet(cached, "WTG data"),
        WTG_FIELDS,
        WTG_REQUIRED,
        _wtg_from_values,
        WTG_COLUMNS,
    )
    issues.extend(wtg_issues)
    programme, programme_issues = _parse_programme(
        sheets["Programme"], _find_sheet(cached, "Programme")
    )
    issues.extend(programme_issues)
    issues.extend(_unknown_endpoints(cables, foundations, wtg_rows))
    if (
        not foundations
        and not cables
        and not wtg_rows
        and not any(issue.code == "missing_columns" for issue in issues)
    ):
        issues.append(
            Issue(
                "error",
                "no_rows",
                "The workbook does not contain any FOU, IAC, or WTG rows.",
            )
        )
    columns: list[str] = []
    for name in [*fou_columns, *iac_columns, *wtg_columns]:
        if name not in columns:
            columns.append(name)
    return ParseResult(
        foundations=foundations,
        cables=cables,
        wtg_rows=wtg_rows,
        programme=programme,
        issues=issues,
        columns=columns,
        source_rows=source_rows,
    )


def _contains_macros(content: bytes) -> bool:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            names = [name.lower() for name in archive.namelist()]
    except zipfile.BadZipFile:
        return False
    return any(name.endswith("vbaproject.bin") for name in names)


def _find_sheet(workbook, expected: str) -> Worksheet | None:
    if workbook is None:
        return None
    target = expected.strip().casefold()
    for name in workbook.sheetnames:
        if name.strip().casefold() == target:
            return workbook[name]
    return None


def _header_key(value: Any, mapping: dict[str, str]) -> str:
    text = str(value or "").replace("\n", " ").strip().casefold()
    text = re.sub(r"\[[^\]]*\]", "", text)
    text = text.replace("aver.", "aver")
    text = re.sub(r"\s+", " ", text).strip(" .")
    return mapping.get(text, "")


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _parse_table(worksheet, cached, mapping, required, build_row, column_names):
    issues: list[Issue] = []
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1))
    fields = [_header_key(cell.value, mapping) for cell in header_row]
    present = {name for name in fields if name}
    missing = [FIELD_LABELS[name] for name in required if name not in present]
    columns = [column_names[name] for name in present if name in column_names]
    if missing:
        issues.append(
            Issue(
                "error",
                "missing_columns",
                f"{worksheet.title} is missing columns: " + ", ".join(missing) + ".",
            )
        )
        return [], [], columns, issues

    rows = []
    source_rows: list[int] = []
    seen: set[str] = set()
    used_cached_formula = False
    for excel_row in worksheet.iter_rows(min_row=2):
        values: dict[str, Any] = {}
        empty = True
        row_number = excel_row[0].row
        formula_error = False
        for cell in excel_row:
            if cell.column - 1 >= len(fields):
                continue
            field_name = fields[cell.column - 1]
            if not field_name:
                continue
            cached_cell = cached.cell(cell.row, cell.column) if cached is not None else None
            value, kind = _read_cell(cell, cached_cell)
            if kind == "formula_missing":
                formula_error = True
                issues.append(
                    Issue(
                        "error",
                        "formula",
                        f"Row {cell.row} column {FIELD_LABELS.get(field_name, field_name)} contains a formula. "
                        "Sequencer does not calculate formulas. Save the workbook from Excel so the value is stored.",
                        row_number=cell.row,
                        column_name=FIELD_LABELS.get(field_name, field_name),
                    )
                )
                continue
            if kind == "cached":
                used_cached_formula = True
            if not _blank(value):
                empty = False
            values[field_name] = value
        if empty:
            continue
        identity = _text(values.get("wtg_id") or values.get("cable_id"))
        if identity.casefold().startswith("note"):
            continue
        if formula_error and identity == "":
            continue
        item, row_issues = build_row(values, row_number, seen)
        issues.extend(row_issues)
        if item is None:
            continue
        rows.append(item)
        source_rows.append(row_number)
        seen.add(_identity(item).casefold())
    if used_cached_formula:
        issues.append(
            Issue(
                "warning",
                "cached_formula",
                "Some cells contain formulas. Sequencer used the values already stored in the workbook and did not calculate them.",
            )
        )
    return rows, source_rows, columns, issues


def _identity(item: Any) -> str:
    if isinstance(item, FoundationInput):
        return item.foundation_id
    if isinstance(item, CableRow):
        return item.cable_id
    return item.wtg_id


def _read_cell(cell: Cell, cached: Cell | None) -> tuple[Any, str]:
    if cell.data_type == "f":
        if cached is None or cached.value is None:
            return None, "formula_missing"
        return cached.value, "cached"
    return cell.value, "value"


def _foundation_from_values(
    values: dict[str, Any], row_number: int, seen: set[str]
) -> tuple[FoundationInput | None, list[Issue]]:
    issues: list[Issue] = []
    foundation_id = _text(values.get("wtg_id"))
    if foundation_id == "":
        issues.append(
            Issue(
                "error",
                "missing_foundation_id",
                f"Row {row_number} is missing WTG ID.",
                row_number=row_number,
                column_name="WTG ID",
            )
        )
        return None, issues
    if foundation_id.casefold() in seen:
        issues.append(
            Issue(
                "error",
                "duplicate_foundation_id",
                f"WTG ID {foundation_id} is duplicated.",
                row_number=row_number,
                column_name="WTG ID",
            )
        )
    include, include_error = _boolean(values.get("include"), required=True)
    if include_error:
        issues.append(
            Issue(
                "error",
                "invalid_boolean",
                f"Row {row_number} Include must be Yes or No.",
                row_number=row_number,
                column_name="Include",
            )
        )
    weight, weight_error = _decimal(values.get("weight"), positive=True, required=include is True)
    if weight_error:
        issues.append(
            Issue(
                "error",
                "invalid_weight",
                f"Row {row_number} Aver. Weight[t] must be a positive number when Include is Yes.",
                row_number=row_number,
                column_name="Aver. Weight[t]",
            )
        )
    easting, east_error = _decimal(values.get("easting"), positive=False, required=True)
    northing, north_error = _decimal(values.get("northing"), positive=False, required=True)
    if east_error or north_error:
        issues.append(
            Issue(
                "error",
                "invalid_coordinate",
                f"Row {row_number} Easting and Northing must be numbers.",
                row_number=row_number,
                column_name="Easting",
            )
        )
    length, length_error = _decimal(values.get("length"), positive=False, required=False)
    if length_error:
        issues.append(
            Issue(
                "error",
                "invalid_number",
                f"Row {row_number} Length[m] must be a number when provided.",
                row_number=row_number,
                column_name="Length[m]",
            )
        )
    eccentricity, eccentricity_error = _decimal(
        values.get("eccentricity_limit"), positive=False, required=False
    )
    if eccentricity_error:
        issues.append(
            Issue(
                "error",
                "invalid_number",
                f"Row {row_number} Eccentricity Limit must be a number when provided.",
                row_number=row_number,
                column_name="Eccentricity Limit",
            )
        )
    spider, spider_error = _optional_boolean(
        values, "spider_location", row_number, "Spider location"
    )
    shallow, shallow_error = _optional_boolean(
        values, "shallow_location", row_number, "Shallow location"
    )
    punch, punch_error = _optional_boolean(values, "punch_through", row_number, "Punch through")
    issues.extend(spider_error + shallow_error + punch_error)
    attributes: dict[str, Any] = {
        "Structure": _text(values.get("structure")),
        "Spider location": spider,
        "Shallow location": shallow,
        "Jacking difficulty": _text(values.get("jacking_difficulty")),
        "Soil difficulty": _text(values.get("soil_difficulty")),
    }
    return (
        FoundationInput(
            foundation_id=foundation_id,
            include=include is True,
            weight_t=weight,
            length_m=length,
            delivery_batch=_text(values.get("delivery_batch")),
            easting_m=easting,
            northing_m=northing,
            punch_through=punch,
            eccentricity_limit_m=eccentricity,
            attributes=attributes,
        ),
        issues,
    )


def _cable_from_values(
    values: dict[str, Any], row_number: int, seen: set[str]
) -> tuple[CableRow | None, list[Issue]]:
    issues: list[Issue] = []
    cable_id = _text(values.get("cable_id"))
    if cable_id == "":
        issues.append(
            Issue(
                "error",
                "missing_cable_id",
                f"Row {row_number} is missing Cable ID.",
                row_number=row_number,
                column_name="Cable ID",
            )
        )
        return None, issues
    if cable_id.casefold() in seen:
        issues.append(
            Issue(
                "error",
                "duplicate_cable_id",
                f"Cable ID {cable_id} is duplicated.",
                row_number=row_number,
                column_name="Cable ID",
            )
        )
    include, include_error = _boolean(values.get("include"), required=True)
    if include_error:
        issues.append(
            Issue(
                "error",
                "invalid_boolean",
                f"Row {row_number} Include must be Yes or No.",
                row_number=row_number,
                column_name="Include",
            )
        )
    spider, spider_issues = _optional_boolean(values, "spider_cable", row_number, "Spider Cable")
    issues.extend(spider_issues)
    return (
        CableRow(
            cable_id=cable_id,
            include=True if include is True else False if include is False else None,
            nickname=_text(values.get("nickname")),
            string_id=_text(values.get("string_id")),
            loop_id=_text(values.get("loop_id")),
            start_location_id=_text(values.get("start_location")),
            end_location_id=_text(values.get("end_location")),
            spider_cable=spider,
            source_row=row_number,
        ),
        issues,
    )


def _wtg_from_values(
    values: dict[str, Any], row_number: int, seen: set[str]
) -> tuple[WtgRow | None, list[Issue]]:
    issues: list[Issue] = []
    wtg_id = _text(values.get("wtg_id"))
    if wtg_id == "":
        issues.append(
            Issue(
                "error",
                "missing_foundation_id",
                f"Row {row_number} is missing WTG ID.",
                row_number=row_number,
                column_name="WTG ID",
            )
        )
        return None, issues
    if wtg_id.casefold() in seen:
        issues.append(
            Issue(
                "error",
                "duplicate_foundation_id",
                f"WTG ID {wtg_id} is duplicated on WTG data.",
                row_number=row_number,
                column_name="WTG ID",
            )
        )
    include, include_error = _boolean(values.get("include"), required=True)
    if include_error:
        issues.append(
            Issue(
                "error",
                "invalid_boolean",
                f"Row {row_number} Include must be Yes or No.",
                row_number=row_number,
                column_name="Include",
            )
        )
    easting, east_error = _decimal(values.get("easting"), positive=False, required=include is True)
    northing, north_error = _decimal(
        values.get("northing"), positive=False, required=include is True
    )
    if east_error or north_error:
        issues.append(
            Issue(
                "error",
                "invalid_coordinate",
                f"Row {row_number} Easting and Northing must be numbers when Include is Yes.",
                row_number=row_number,
                column_name="Easting",
            )
        )
    spider, spider_issues = _optional_boolean(
        values, "spider_location", row_number, "Spider location"
    )
    issues.extend(spider_issues)
    return (
        WtgRow(
            wtg_id=wtg_id,
            include=True if include is True else False if include is False else None,
            structure=_text(values.get("structure")),
            easting_m=easting,
            northing_m=northing,
            spider_location=spider,
            string_label=_text(values.get("string")),
            jacking_difficulty=_text(values.get("jacking_difficulty")),
            source_row=row_number,
        ),
        issues,
    )


def _parse_programme(
    worksheet: Worksheet, cached: Worksheet | None
) -> tuple[list[ProgrammeRow], list[Issue]]:
    issues: list[Issue] = []
    rows = list(worksheet.iter_rows(min_row=1, max_row=2))
    if len(rows) < 2:
        return [], [
            Issue(
                "error",
                "missing_columns",
                "Programme needs a package row and a P50 header row.",
            )
        ]
    packages: list[str] = []
    current = ""
    for cell in rows[0]:
        text = _text(cell.value)
        if text:
            current = _ACTIVITY_HEADERS.get(_plain_header(text), "")
        packages.append(current)
    headers = [_plain_header(cell.value) for cell in rows[1]]
    location_index = next((index for index, name in enumerate(headers) if name == "location"), None)
    dates: dict[tuple[str, str], int] = {}
    for index, name in enumerate(headers):
        package = packages[index] if index < len(packages) else ""
        if package and name in {"p50 start", "p50 end"}:
            kind = "start" if name.endswith("start") else "end"
            dates[(package, kind)] = index
    expected = {
        (activity, kind) for activity in _REQUIRED_ACTIVITIES for kind in ("start", "end")
    }
    present = {activity for activity, _kind in dates}
    incomplete = [
        _ACTIVITY_NAMES[activity]
        for activity in present
        if (activity, "start") not in dates or (activity, "end") not in dates
    ]
    if incomplete:
        return [], [
            Issue(
                "error",
                "missing_columns",
                "Programme must list P50 start and end for " + ", ".join(incomplete) + ".",
            )
        ]
    if location_index is None or expected - set(dates):
        return [], [
            Issue(
                "error",
                "missing_columns",
                "Programme must list Location and P50 start and end for FOU, IAC, and WTG.",
            )
        ]
    parsed: list[ProgrammeRow] = []
    seen: set[int] = set()
    cached_sheet = cached
    for excel_row in worksheet.iter_rows(min_row=3):
        if all(_blank(cell.value) for cell in excel_row):
            continue
        row_number = excel_row[0].row
        slot_value, _kind = _cached_value(excel_row, location_index, cached_sheet)
        slot, slot_error = _optional_int(slot_value)
        if slot_error or slot is None:
            issues.append(
                Issue(
                    "error",
                    "invalid_slot",
                    f"Programme row {row_number} Location must be a positive integer.",
                    row_number=row_number,
                    column_name="Location",
                )
            )
            continue
        if slot in seen:
            issues.append(
                Issue(
                    "error",
                    "duplicate_slot",
                    f"Programme location {slot} is duplicated.",
                    row_number=row_number,
                    column_name="Location",
                )
            )
        seen.add(slot)
        values: dict[str, datetime | None] = {}
        failed = False
        for package, kind in sorted(dates):
            raw, kind_name = _cached_value(excel_row, dates[(package, kind)], cached_sheet)
            if kind_name == "formula_missing":
                failed = True
                issues.append(
                    Issue(
                        "error",
                        "formula",
                        f"Programme row {row_number} contains a formula. Sequencer does not calculate formulas.",
                        row_number=row_number,
                        column_name="P50",
                    )
                )
                continue
            parsed_date, date_error = _datetime(raw)
            if date_error:
                failed = True
                issues.append(
                    Issue(
                        "error",
                        "invalid_date",
                        f"Programme row {row_number} {_ACTIVITY_NAMES.get(package, package)} "
                        f"P50 {kind} must be a date.",
                        row_number=row_number,
                        column_name="P50",
                    )
                )
            values[f"{package}_p50_{kind}"] = parsed_date
        if failed:
            continue
        parsed.append(ProgrammeRow(slot=slot, source_row=row_number, **values))
    return parsed, issues


def _cached_value(excel_row, index: int, cached: Worksheet | None) -> tuple[Any, str]:
    if index >= len(excel_row):
        return None, "value"
    cell = excel_row[index]
    cached_cell = cached.cell(cell.row, cell.column) if cached is not None else None
    return _read_cell(cell, cached_cell)


def _unknown_endpoints(
    cables: list[CableRow], foundations: list[FoundationInput], wtg_rows: list[WtgRow]
) -> list[Issue]:
    known = {item.foundation_id.casefold() for item in foundations}
    known.update(item.wtg_id.casefold() for item in wtg_rows)
    issues: list[Issue] = []
    for cable in cables:
        for label, value in (
            ("start", cable.start_location_id),
            ("end", cable.end_location_id),
        ):
            if value and value.casefold() not in known:
                issues.append(
                    Issue(
                        "warning",
                        "unknown_endpoint",
                        f"Cable {cable.cable_id} {label} location {value} is not a FOU or WTG id.",
                        row_number=cable.source_row,
                        column_name="Location ID",
                    )
                )
    return issues


def _optional_boolean(
    values: dict[str, Any], key: str, row_number: int, label: str
) -> tuple[bool | None, list[Issue]]:
    if key not in values:
        return None, []
    value, error = _boolean(values.get(key), required=False)
    if error:
        return None, [
            Issue(
                "error",
                "invalid_boolean",
                f"Row {row_number} {label} must be Yes or No when provided.",
                row_number=row_number,
                column_name=label,
            )
        ]
    return value, []


def _plain_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().casefold())


def _text(value: Any) -> str:
    if _blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return str(value).strip()


def _boolean(value: Any, *, required: bool) -> tuple[bool | None, bool]:
    if _blank(value):
        return None, required
    if isinstance(value, bool):
        return value, False
    text = _text(value).casefold()
    if text in TRUE_TEXT:
        return True, False
    if text in FALSE_TEXT:
        return False, False
    return None, True


def _decimal(value: Any, *, positive: bool, required: bool) -> tuple[Decimal | None, bool]:
    if _blank(value) or isinstance(value, bool):
        return None, required
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None, True
    if not number.is_finite():
        return None, True
    if positive and number <= 0:
        return None, True
    return number.quantize(Decimal("0.001")), False


def _optional_int(value: Any) -> tuple[int | None, bool]:
    if _blank(value):
        return None, False
    if isinstance(value, bool):
        return None, True
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None, True
    if number != number.to_integral_value() or number < 1:
        return None, True
    return int(number), False


def _datetime(value: Any) -> tuple[datetime | None, bool]:
    if _blank(value):
        return None, False
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime(value.year, value.month, value.day)
    else:
        return None, True
    if timezone.is_naive(parsed):
        parsed = timezone.make_aware(parsed, UTC)
    return parsed, False
