"""Downloadable workbook with FOU, IAC, WTG, and Programme headings.

There is no project location data. Constraints and settings are created in the app.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FOU_HEADERS = (
    "WTG ID",
    "Include",
    "Structure",
    "Easting",
    "Northing",
    "Aver. Weight[t]",
    "Length[m]",
    "Delivery Batch",
    "Spider location",
    "Shallow location",
    "Soil difficulty",
    "Jacking Difficulty",
    "Punch through",
    "Eccentricity Limit",
)
IAC_HEADERS = (
    "Cable ID",
    "Nickname",
    "Include",
    "String ID",
    "Loop ID",
    "Location ID\n(Cable start-point)",
    "Location ID\n(Cable end-point)",
    "Spider Cable",
)
WTG_HEADERS = (
    "WTG ID",
    "Include",
    "Structure",
    "Easting",
    "Northing",
    "Jacking difficulty WTG",
)


def build_template() -> bytes:
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="0F2C4C")
    header_font = Font(color="FFFFFF", bold=True)
    fou = workbook.active
    fou.title = "FOU data"
    _write_header(fou, FOU_HEADERS, header_fill, header_font)
    iac = workbook.create_sheet("IAC data")
    _write_header(iac, IAC_HEADERS, header_fill, header_font)
    for column in (6, 7):
        iac.cell(1, column).alignment = Alignment(wrap_text=True)
    wtg = workbook.create_sheet("WTG data")
    _write_header(wtg, WTG_HEADERS, header_fill, header_font)
    programme = workbook.create_sheet("Programme")
    activities = (
        "FOU installation",
        "SSt installation",
        "Filter layer",
        "Armour layer",
        "IAC lay",
        "IAC burial",
        "WTG installation",
        "WTG commissioning",
    )
    programme.cell(2, 1, "Location")
    for index, name in enumerate(activities):
        column = 2 + index * 2
        programme.cell(1, column, name)
        programme.merge_cells(
            start_row=1, start_column=column, end_row=1, end_column=column + 1
        )
        programme.cell(2, column, "P50 Start")
        programme.cell(2, column + 1, "P50 End")
    for index in range(1, 18):
        programme.column_dimensions[get_column_letter(index)].width = 22
    for row in programme.iter_rows(min_row=1, max_row=2, max_col=17):
        for cell in row:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
    programme.freeze_panes = "A3"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_header(worksheet, headers, fill, font) -> None:
    for index, name in enumerate(headers, start=1):
        cell = worksheet.cell(1, index, name)
        cell.fill = fill
        cell.font = font
        worksheet.column_dimensions[get_column_letter(index)].width = max(18, len(name) + 2)
    worksheet.freeze_panes = "A2"
